"""Чтение прикреплённых при создании агента файлов в текстовый контекст.

Файл копируется в рабочую папку агента (чтобы Excel-инструменты могли с ним
работать), а его содержимое кратко извлекается в текст для передачи в LLM.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from agent_desktop_constructor.tools.agent_workspace import AgentWorkspace

_TEXT_SUFFIXES = {".txt", ".csv", ".md", ".json", ".log", ".tsv", ".yaml", ".yml"}
_MAX_TEXT_CHARS = 8000
_MAX_XLSX_ROWS = 100


def ingest_attachment(workspace: AgentWorkspace, source_path: str) -> dict:
    """Скопировать файл в папку агента и вернуть {name, content}.

    Никогда не бросает исключение наружу: при ошибке возвращает пояснение в
    поле content, чтобы поток создания агента не прерывался.
    """
    source = Path(source_path)
    name = source.name
    try:
        target = workspace.resolve(name)
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)
    except Exception as exc:  # noqa: BLE001 - файл-вложение не должен ронять запуск
        return {"name": name, "content": f"(не удалось скопировать файл: {exc})"}

    return {"name": name, "content": _summarize(target)}


def _summarize(path: Path) -> str:
    """Извлечь краткое текстовое представление файла для контекста LLM."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return _summarize_xlsx(path)
        if suffix in _TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="replace")
            if len(text) > _MAX_TEXT_CHARS:
                return text[:_MAX_TEXT_CHARS] + "\n… (текст обрезан)"
            return text
    except Exception as exc:  # noqa: BLE001
        return f"(не удалось прочитать содержимое: {exc})"
    size = path.stat().st_size if path.exists() else 0
    return (
        f"(файл {path.name}, {size} байт, скопирован в рабочую папку агента; "
        "используйте инструменты для его обработки)"
    )


def _summarize_xlsx(path: Path) -> str:
    """Сжать содержимое .xlsx в текст: листы и первые строки."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = [f"Excel-файл, листы: {list(workbook.sheetnames)}"]
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            lines.append(f"[Лист {sheet_name}]")
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= _MAX_XLSX_ROWS:
                    lines.append("… (строки обрезаны)")
                    break
                cells = [
                    "" if cell is None else str(cell)
                    for cell in row
                ]
                lines.append(" | ".join(cells))
        return "\n".join(lines)
    finally:
        workbook.close()
