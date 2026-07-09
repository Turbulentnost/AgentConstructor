import re
import json
import os

try:
    from bs4 import BeautifulSoup
except ImportError:
    print("ERROR: BeautifulSoup not available")
    exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    print("ERROR: openpyxl not available")
    exit(1)

# Read the HTML file
with open('page_dumps/roseltorg_procedures_20260709_091227/page.html', 'r', encoding='utf-8') as f:
    html_content = f.read()

soup = BeautifulSoup(html_content, 'html.parser')

# Find grid rows
rows = soup.find_all('div', class_=re.compile(r'x-grid3-row'))
print(f"Total grid rows found: {len(rows)}")

# Parse rows and filter by status
procedures = []
for row in rows:
    cells = row.find_all('td')
    if not cells:
        continue
    cell_texts = [cell.get_text(strip=True) for cell in cells]
    
    # Skip rows that don't have enough cells or no status
    if len(cell_texts) < 18:
        continue
    
    status = cell_texts[17]
    if status != 'Приём заявок':
        continue
    
    # Extract deadline text - clean up 'Осталось X дней' part
    deadline_raw = cell_texts[13]
    deadline_match = re.match(r'(\d{2}\.\d{2}\.\d{4}\s+\d{2}:\d{2})', deadline_raw)
    deadline = deadline_match.group(1) if deadline_match else deadline_raw
    
    # Extract sum - clean up
    summa = cell_texts[16] if cell_texts[16] != '—' else ''
    
    procedure = {
        'registry_number': cell_texts[2],
        'procedure_number': cell_texts[4],
        'joint': cell_texts[6],
        'organizer': cell_texts[7],
        'contact_person': cell_texts[8],
        'name': cell_texts[9],
        'customer': cell_texts[10],
        'publish_date': cell_texts[11],
        'deadline': deadline,
        'sum': summa,
        'status': status
    }
    procedures.append(procedure)

# Take only first 25 (first page)
procedures = procedures[:25]
print(f"Procedures with status 'Приём заявок' on first page: {len(procedures)}")

# Create Excel workbook
wb = Workbook()
ws = wb.active
ws.title = 'Процедуры - Приём заявок'

# Headers
headers = ['№', 'Реестровый №', 'Номер процедуры', 'Совместная', 'Организатор', 
           'Контактное лицо', 'Наименование', 'Заказчик', 'Дата публикации', 
           'Приём заявок до', 'Сумма', 'Статус']

# Style for headers
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font_white = Font(bold=True, size=11, color='FFFFFF')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write headers
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

# Write data
for idx, proc in enumerate(procedures, 1):
    row_num = idx + 1
    ws.cell(row=row_num, column=1, value=idx).border = thin_border
    ws.cell(row=row_num, column=2, value=proc['registry_number']).border = thin_border
    ws.cell(row=row_num, column=3, value=proc['procedure_number']).border = thin_border
    ws.cell(row=row_num, column=4, value=proc['joint']).border = thin_border
    ws.cell(row=row_num, column=5, value=proc['organizer']).border = thin_border
    ws.cell(row=row_num, column=6, value=proc['contact_person']).border = thin_border
    ws.cell(row=row_num, column=7, value=proc['name']).border = thin_border
    ws.cell(row=row_num, column=8, value=proc['customer']).border = thin_border
    ws.cell(row=row_num, column=9, value=proc['publish_date']).border = thin_border
    ws.cell(row=row_num, column=10, value=proc['deadline']).border = thin_border
    ws.cell(row=row_num, column=11, value=proc['sum']).border = thin_border
    ws.cell(row=row_num, column=12, value=proc['status']).border = thin_border
    
    # Wrap text for long fields
    ws.cell(row=row_num, column=5).alignment = Alignment(wrap_text=True)
    ws.cell(row=row_num, column=7).alignment = Alignment(wrap_text=True)
    ws.cell(row=row_num, column=8).alignment = Alignment(wrap_text=True)

# Set column widths
column_widths = [5, 12, 20, 10, 30, 25, 50, 30, 14, 20, 20, 15]
for col, width in enumerate(column_widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

# Save
output_path = 'roseltorg_procedures_priem_zayavok.xlsx'
wb.save(output_path)
print(f"\nExcel file saved: {output_path}")
print(f"Total rows written: {len(procedures)}")

# Print summary of first 5 procedures
print("\n--- First 5 procedures ---")
for i, p in enumerate(procedures[:5], 1):
    print(f"{i}. {p['procedure_number']} | {p['organizer'][:30]} | {p['name'][:50]} | {p['deadline']}")
