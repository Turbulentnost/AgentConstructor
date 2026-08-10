import openpyxl, json, os

files = [
    'ГРАФИК ОТГРУЗОК (расширенный) 2.xlsx',
    'График производства 1.xlsx',
    'детальный.xlsx',
    'ТАМОЖНЯ.xlsx',
]

os.makedirs('code/dump', exist_ok=True)

for fname in files:
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
    except Exception as e:
        print(f'ERROR opening {fname}: {e}')
        continue
    result = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        max_row = min(ws.max_row, 400)
        max_col = min(ws.max_column, 60)
        for r in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            row_vals = []
            for c in r:
                v = c.value
                if hasattr(v, 'isoformat'):
                    v = v.isoformat()
                row_vals.append(v)
            # skip fully empty rows at end but keep for now
            rows.append(row_vals)
        result[sheet_name] = {
            'max_row': ws.max_row,
            'max_col': ws.max_column,
            'rows': rows,
        }
    out_name = 'code/dump/' + fname.replace('.xlsx', '').replace(' ', '_').replace('(', '').replace(')', '') + '.json'
    with open(out_name, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False)
    print(f'Saved {fname} -> {out_name}, sheets: {wb.sheetnames}')
