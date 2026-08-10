import openpyxl
import os

files = [
    "ГРАФИК ОТГРУЗОК (расширенный) 2.xlsx",
    "График производства 1.xlsx",
    "детальный.xlsx",
    "ТАМОЖНЯ.xlsx",
]

os.makedirs("reports", exist_ok=True)
out_path = os.path.join("reports", "structure_report.txt")

with open(out_path, "w", encoding="utf-8") as out:
    for fname in files:
        if not os.path.exists(fname):
            out.write(f"=== FILE NOT FOUND: {fname} ===\n")
            continue
        out.write(f"\n{'='*100}\nFILE: {fname}\n{'='*100}\n")
        try:
            wb = openpyxl.load_workbook(fname, data_only=True)
        except Exception as e:
            out.write(f"  ERROR loading: {e}\n")
            continue
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            out.write(f"\n--- SHEET: '{sheet_name}' | dims: {ws.dimensions} | max_row={ws.max_row} max_col={ws.max_column} ---\n")
            max_r = min(ws.max_row, 10)
            max_c = min(ws.max_column, 20)
            for r in range(1, max_r + 1):
                row_vals = []
                for c in range(1, max_c + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, str) and len(v) > 35:
                        v = v[:35] + "..."
                    row_vals.append(str(v) if v is not None else "")
                out.write(f"  row{r}: {row_vals}\n")
            # Search for keywords indicating stock/specification/prices roles
            keywords = ["остаток", "остатки", "спецификация", "bom", "цена", "поставщик", "ед. изм", "единица измерения", "номенклатура материала"]
            found_kw = set()
            for r in range(1, min(ws.max_row, 15) + 1):
                for c in range(1, min(ws.max_column, 30) + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, str):
                        low = v.lower()
                        for kw in keywords:
                            if kw in low:
                                found_kw.add(kw)
            if found_kw:
                out.write(f"  >>> KEYWORDS FOUND: {sorted(found_kw)}\n")

print(f"Report written to {out_path}")
print(f"File size: {os.path.getsize(out_path)} bytes")
