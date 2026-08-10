import openpyxl
import os

files = [
    "ГРАФИК ОТГРУЗОК (расширенный) 2.xlsx",
    "График производства 1.xlsx",
    "детальный.xlsx",
    "ТАМОЖНЯ.xlsx",
]

for fname in files:
    if not os.path.exists(fname):
        print(f"=== FILE NOT FOUND: {fname} ===")
        continue
    print(f"\n{'='*80}\nFILE: {fname}\n{'='*80}")
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
    except Exception as e:
        print(f"  ERROR loading: {e}")
        continue
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- SHEET: '{sheet_name}' | dims: {ws.dimensions} | max_row={ws.max_row} max_col={ws.max_column} ---")
        # print first up to 6 rows, up to 15 columns, to see header structure
        max_r = min(ws.max_row, 6)
        max_c = min(ws.max_column, 15)
        for r in range(1, max_r + 1):
            row_vals = []
            for c in range(1, max_c + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and len(v) > 40:
                    v = v[:40] + "..."
                row_vals.append(str(v) if v is not None else "")
            print(f"  row{r}: {row_vals}")
print("\nDONE")
