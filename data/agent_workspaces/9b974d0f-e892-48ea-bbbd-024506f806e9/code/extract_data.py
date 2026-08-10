# -*- coding: utf-8 -*-
"""
Extraction-скрипт: читает все 4 .xlsx рабочей папки и сохраняет структурированные
данные в JSON-файлы для дальнейшего расчёта. Печатает в stdout только короткое
ASCII-safe summary (без построчных кириллических дампов) чтобы избежать mojibake.
"""
import os
import glob
import json
import datetime

from openpyxl import load_workbook


def cell(ws, r, c):
    return ws.cell(row=r, column=c).value


def dump_sheet_raw(ws, max_rows=None, max_cols=None):
    max_rows = max_rows or ws.max_row
    max_cols = max_cols or ws.max_column
    rows = []
    for r in range(1, max_rows + 1):
        row = []
        for c in range(1, max_cols + 1):
            v = cell(ws, r, c)
            if isinstance(v, (datetime.date, datetime.datetime)):
                v = v.isoformat()
            row.append(v)
        rows.append(row)
    return rows


def main():
    cwd = os.getcwd()
    files = sorted(glob.glob("*.xlsx"))
    result = {"cwd": cwd, "files": {}}

    for fn in files:
        try:
            wb = load_workbook(fn, data_only=True, read_only=False)
        except Exception as e:
            result["files"][fn] = {"error": repr(e)}
            continue
        f_data = {"sheets": {}}
        for ws in wb.worksheets:
            # ограничим объём: только реально используемые строки/колонки
            max_r = min(ws.max_row or 0, 500)
            max_c = min(ws.max_column or 0, 80)
            rows = dump_sheet_raw(ws, max_r, max_c)
            f_data["sheets"][ws.title] = {
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "rows": rows,
            }
        wb.close()
        result["files"][fn] = f_data

    out_path = "raw_extracted_data.json"
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(result, fh, ensure_ascii=False, indent=1)

    # ASCII-safe summary без кириллицы построчно
    print("OK. Saved raw data to:", out_path)
    print("Files processed:", len(files))
    for fn in files:
        fd = result["files"].get(fn, {})
        if "error" in fd:
            print("  [ERROR]", fn.encode("ascii", "backslashreplace").decode("ascii"))
            continue
        sheets = fd.get("sheets", {})
        print("  file idx ok, sheets_count=", len(sheets))
        for sname, sdata in sheets.items():
            print("    sheet rows=%s cols=%s" % (sdata["max_row"], sdata["max_col"]))


if __name__ == "__main__":
    main()
