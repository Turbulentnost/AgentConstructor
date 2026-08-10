# -*- coding: utf-8 -*-
"""Инспектор структуры всех .xlsx в рабочей папке агента.
Сохраняет подробный отчёт в inspect_report.txt, в stdout печатает краткое summary.
"""
import os
import glob
import json

from openpyxl import load_workbook

MAX_ROWS = 18          # сколько первых строк показывать по листу
MAX_COLS = 28          # сколько первых колонок показывать
CELL_LIMIT = 40        # обрезка длинных значений

MARKERS = [
    "номенклатур", "остаток", "ед. изм", "ед.изм", "цена", "поставщик",
    "п/ф", "отк", "склад", "версия", "заказ", "опытн", "наименование",
    "таможн", "мск", "ростов", "итц", "реестр", "дата постав", "спецификац",
    "изделие", "модель", "кол-во", "артикул", "партия", "недел", "итог",
]


def fmt(v):
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    if len(s) > CELL_LIMIT:
        s = s[: CELL_LIMIT - 1] + "~"
    return s


def inspect_sheet(ws):
    info = {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "preview": [],
        "markers": set(),
        "nonempty_rows": 0,
    }
    limit_rows = min(ws.max_row or 0, 400)
    limit_cols = min(ws.max_column or 0, 80)
    for r in range(1, limit_rows + 1):
        row_vals = []
        any_val = False
        for c in range(1, limit_cols + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                any_val = True
                low = str(v).lower()
                for m in MARKERS:
                    if m in low:
                        info["markers"].add(m)
            if c <= MAX_COLS:
                row_vals.append(fmt(v))
        if any_val:
            info["nonempty_rows"] += 1
        if r <= MAX_ROWS:
            while row_vals and row_vals[-1] == "":
                row_vals.pop()
            info["preview"].append(row_vals)
    info["markers"] = sorted(info["markers"])
    return info


def main():
    cwd = os.getcwd()
    files = sorted(glob.glob("*.xlsx"))
    report_lines = []
    summary = []
    report_lines.append("CWD: %s" % cwd)
    report_lines.append("FILES: %s" % files)

    for fn in files:
        report_lines.append("\n" + "=" * 90)
        report_lines.append("FILE: %s (%d bytes)" % (fn, os.path.getsize(fn)))
        try:
            wb = load_workbook(fn, data_only=True, read_only=False)
        except Exception as e:
            report_lines.append("  !! ошибка открытия: %r" % (e,))
            summary.append({"file": fn, "error": repr(e)})
            continue
        f_summary = {"file": fn, "sheets": []}
        report_lines.append("SHEETS: %s" % wb.sheetnames)
        for ws in wb.worksheets:
            info = inspect_sheet(ws)
            report_lines.append("\n  --- SHEET '%s'  rows=%s cols=%s nonempty=%s" % (
                info["title"], info["max_row"], info["max_col"], info["nonempty_rows"]))
            report_lines.append("      markers: %s" % ", ".join(info["markers"]))
            for i, row in enumerate(info["preview"], start=1):
                report_lines.append("      r%-3d | %s" % (i, " | ".join(row)))
            f_summary["sheets"].append({
                "sheet": info["title"],
                "rows": info["max_row"],
                "cols": info["max_col"],
                "nonempty": info["nonempty_rows"],
                "markers": info["markers"],
            })
        wb.close()
        summary.append(f_summary)

    with open("inspect_report.txt", "w", encoding="utf-8") as fh:
        fh.write("\n".join(report_lines))

    with open("inspect_summary.json", "w", encoding="utf-8") as fh:
        json.dump(summary, fh, ensure_ascii=False, indent=1)

    # компактный stdout
    print("REPORT: inspect_report.txt  SUMMARY: inspect_summary.json")
    for f in summary:
        if "error" in f:
            print("FILE %s -> ERROR %s" % (f["file"], f["error"]))
            continue
        print("\nFILE %s" % f["file"])
        for s in f["sheets"]:
            print("  sheet='%s' rows=%s cols=%s nonempty=%s markers=%s" % (
                s["sheet"], s["rows"], s["cols"], s["nonempty"], ",".join(s["markers"][:12])))


if __name__ == "__main__":
    main()
