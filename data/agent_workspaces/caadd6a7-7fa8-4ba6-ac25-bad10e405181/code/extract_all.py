import json, os
import openpyxl

files = [
    "ГРАФИК ОТГРУЗОК (расширенный) 2.xlsx",
    "График производства 1.xlsx",
    "детальный.xlsx",
    "ТАМОЖНЯ.xlsx",
]

base_dir = os.getcwd()
# Файлы лежат в родительской папке агента (code/ - подпапка)
candidates = [base_dir, os.path.dirname(base_dir)]

def find_file(name):
    for d in candidates:
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    return None

result = {}
for fname in files:
    path = find_file(fname)
    if not path:
        result[fname] = {"error": "not found", "tried": candidates}
        continue
    wb = openpyxl.load_workbook(path, data_only=True)
    file_data = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            # convert non-serializable (datetime) to str
            row2 = []
            for v in row:
                if hasattr(v, 'isoformat'):
                    row2.append(v.isoformat())
                else:
                    row2.append(v)
            rows.append(row2)
        file_data[sheet_name] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "rows": rows,
        }
    result[fname] = file_data

out_path = os.path.join(base_dir, "data_dump.json")
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

print("Saved to", out_path)
for fname, fdata in result.items():
    if isinstance(fdata, dict) and "error" in fdata:
        print(fname, "ERROR", fdata["error"])
        continue
    print(fname, "sheets:", list(fdata.keys()))
    for sn, sd in fdata.items():
        print("  ", sn, sd["max_row"], sd["max_col"])
